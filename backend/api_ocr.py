from fastapi import FastAPI, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, StreamingResponse, Response
from pydantic import BaseModel
from typing import List, Optional
import base64
import cv2
import json
import numpy as np
import os
import time
import io

import httpx
from openai import OpenAI
from paddleocr import PaddleOCR
from pipeline import parse_elements, check_all_services, layout_analyze

app = FastAPI(title="Document Parsing Services")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

ocr_pipeline = None

def get_ocr():
    global ocr_pipeline
    if ocr_pipeline is None:
        os.environ["PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK"] = "True"
        print("Initializing OCR (PP-OCRv5)...")
        ocr_pipeline = PaddleOCR(ocr_version="PP-OCRv5", lang="ch")
    return ocr_pipeline

class LayoutBBox(BaseModel):
    poly: List[float]
    category_type: str

class OCRParseRequest(BaseModel):
    image_base64: str
    layout_bboxes: List[LayoutBBox]
    merge_text: bool = True

class LayoutRequest(BaseModel):
    image_base64: str
    threshold: float = 0.50
    layout_nms: bool = True

@app.middleware("http")
async def log_requests(request: Request, call_next):
    start = time.time()
    response = await call_next(request)
    duration_ms = round((time.time() - start) * 1000)
    path = request.url.path
    if path == "/api/health":
        return response
    print(json.dumps({
        "method": request.method,
        "path": path,
        "status": response.status_code,
        "duration_ms": duration_ms,
    }))
    return response

@app.exception_handler(Exception)
async def global_exception_handler(request: Request, exc: Exception):
    path = request.url.path
    error_msg = str(exc)[:500]
    print(json.dumps({
        "method": request.method,
        "path": path,
        "status": 500,
        "duration_ms": 0,
        "error": error_msg,
    }))
    return JSONResponse(status_code=500, content={"detail": str(exc)})

@app.get("/api/health")
async def health_check():
    ocr = get_ocr()
    ocr_status = "loaded" if ocr is not None else "loading"
    
    remote_status = {}
    try:
        report = await check_all_services()
        for s in report.services:
            remote_status[s.name] = {
                "available": s.available,
                "latency_ms": s.latency_ms,
            }
    except Exception:
        remote_status["error"] = "unable to check remote services"

    return JSONResponse(
        status_code=200,
        content={
            "status": "ok",
            "models": {"ocr": ocr_status, "layout": "remote (doclayout:8765)"},
            "remote_services": remote_status,
        },
    )

@app.post("/api/layout")
async def detect_layout(req: LayoutRequest):
    try:
        result = await layout_analyze(req.image_base64)
        return result
    except Exception as e:
        raise

@app.post("/api/parse")
async def parse_elements_endpoint(req: OCRParseRequest):
    try:
        img_data = req.image_base64.split(",")[-1] if "," in req.image_base64 else req.image_base64
        img_bytes = base64.b64decode(img_data)
        img = cv2.imdecode(np.frombuffer(img_bytes, np.uint8), cv2.IMREAD_COLOR)

        if img is None:
            raise HTTPException(status_code=400, detail="Invalid image data")

        results = await parse_elements(
            image_base64=req.image_base64,
            layout_bboxes=req.layout_bboxes,
            merge_text=req.merge_text,
        )
        return {"elements": results}
    except Exception as e:
        raise

# ── AI Companion Reading ──────────────────────────────────────

DEEPSEEK_API_KEY = os.environ.get("DEEPSEEK_API_KEY", "sk-558df2e5f3dd4592b498128243c032ef")
DEEPSEEK_MODEL = os.environ.get("DEEPSEEK_MODEL", "deepseek-v4-pro")
DEEPSEEK_BASE_URL = "https://api.deepseek.com"

_client: Optional[OpenAI] = None

def _get_client() -> OpenAI:
    global _client
    if _client is None:
        _client = OpenAI(api_key=DEEPSEEK_API_KEY, base_url=DEEPSEEK_BASE_URL)
    return _client

AI_SYSTEM_PROMPTS = {
    "translate": "You are a professional translator. Translate the following text accurately while preserving the meaning, tone, and formatting. Output only the translated text without any explanations.",
    "explain": "You are a knowledgeable assistant. Explain the following text clearly and concisely. Use markdown formatting for readability. Keep explanations focused and avoid unnecessary tangents.",
    "summarize": "You are a summarization expert. Summarize the following content concisely, capturing all key points. Use markdown formatting with bullet points where appropriate.",
    "chat": "You are a helpful document assistant. Answer questions about the document content provided in context. Be concise and accurate. Use markdown formatting.",
}


class AIActionRequest(BaseModel):
    action: str
    content: str
    params: dict = {}


class AIActionResponse(BaseModel):
    result: str


@app.post("/api/ai/action", response_model=AIActionResponse)
async def ai_action(req: AIActionRequest):
    try:
        system_prompt = AI_SYSTEM_PROMPTS.get(req.action, AI_SYSTEM_PROMPTS["chat"])
        target_lang = req.params.get("target_lang", "Chinese")

        if req.action == "translate":
            user_prompt = f"Translate the following text to {target_lang}:\n\n{req.content}"
        elif req.action == "explain":
            user_prompt = f"Please explain the following text:\n\n{req.content}"
        elif req.action == "summarize":
            user_prompt = f"Please summarize the following content:\n\n{req.content}"
        elif req.action == "chat":
            user_prompt = req.content
        else:
            raise HTTPException(status_code=400, detail=f"Unknown action: {req.action}")

        client = _get_client()
        resp = client.chat.completions.create(
            model=DEEPSEEK_MODEL,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            temperature=0.3,
            max_tokens=4096,
        )
        result = resp.choices[0].message.content or ""
        return AIActionResponse(result=result)
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/ai/stream")
async def ai_stream(req: AIActionRequest):
    system_prompt = AI_SYSTEM_PROMPTS.get(req.action, AI_SYSTEM_PROMPTS["chat"])
    target_lang = req.params.get("target_lang", "Chinese")

    if req.action == "translate":
        user_prompt = f"Translate the following text to {target_lang}:\n\n{req.content}"
    elif req.action == "explain":
        user_prompt = f"Please explain the following text:\n\n{req.content}"
    elif req.action == "summarize":
        user_prompt = f"Please summarize the following content:\n\n{req.content}"
    else:
        user_prompt = req.content

    async def generate():
        try:
            client = _get_client()
            stream = client.chat.completions.create(
                model=DEEPSEEK_MODEL,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_prompt},
                ],
                temperature=0.3,
                max_tokens=4096,
                stream=True,
            )
            for chunk in stream:
                delta = chunk.choices[0].delta
                if delta.content:
                    yield f"data: {json.dumps({'content': delta.content})}\n\n"
            yield "data: [DONE]\n\n"
        except Exception as e:
            yield f"data: {json.dumps({'error': str(e)})}\n\n"

    return StreamingResponse(generate(), media_type="text/event-stream")


# ── Format Conversion ─────────────────────────────────────────

class ConvertRequest(BaseModel):
    format: str
    markdown: str
    title: str = "Document"


@app.post("/api/convert")
async def convert_document(req: ConvertRequest):
    try:
        if req.format == "html":
            html = markdown_to_html(req.markdown, req.title)
            return Response(
                content=html,
                media_type="text/html",
                headers={"Content-Disposition": f"attachment; filename={req.title}.html"},
            )
        elif req.format == "docx":
            docx_bytes = markdown_to_docx(req.markdown, req.title)
            return StreamingResponse(
                io.BytesIO(docx_bytes),
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={"Content-Disposition": f"attachment; filename={req.title}.docx"},
            )
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported format: {req.format}")
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


def markdown_to_html(md_text: str, title: str) -> str:
    import markdown as md_lib
    body = md_lib.markdown(md_text, extensions=["tables", "fenced_code", "codehilite"])
    return f"""<!DOCTYPE html>
<html lang="zh">
<head><meta charset="utf-8"><title>{title}</title>
<style>
body {{ font-family: -apple-system, sans-serif; max-width: 800px; margin: 40px auto; padding: 0 20px; line-height: 1.8; }}
table {{ border-collapse: collapse; width: 100%; margin: 16px 0; }}
th, td {{ border: 1px solid #ddd; padding: 8px 12px; text-align: left; }}
th {{ background: #f5f5f5; }}
img {{ max-width: 100%; }}
pre {{ background: #f5f5f5; padding: 16px; border-radius: 6px; overflow-x: auto; }}
</style></head>
<body>{body}</body></html>"""


def markdown_to_docx(md_text: str, title: str) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt, Inches

        doc = Document()
        doc.styles["Normal"].font.size = Pt(11)
        doc.add_heading(title, level=0)

        for line in md_text.split("\n"):
            line = line.strip()
            if not line:
                continue
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("|"):
                continue  # skip table lines (complex parsing needed)
            else:
                doc.add_paragraph(line)

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
    except ImportError:
        raise HTTPException(status_code=503, detail="python-docx not installed")


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8002)

